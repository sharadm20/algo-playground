from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

def create_day1_document():
    # Create a new document
    doc = Document()
    
    # Title
    title = doc.add_heading('30-Day DSA Study Plan - Day 1: Arrays & Hashing', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Introduction
    doc.add_heading('Topic Overview', level=1)
    doc.add_paragraph(
        'Arrays and Hashing form the foundation of many algorithmic solutions. '
        'Understanding these data structures and their associated techniques is crucial '
        'for solving a wide variety of problems efficiently.'
    )
    
    # Conceptual Foundation
    doc.add_heading('Conceptual Foundation', level=1)
    doc.add_paragraph(
        '• Arrays: Contiguous memory allocation allowing O(1) random access\n'
        '• Hash Tables: Key-value mappings with average O(1) lookup using hash functions\n'
        '• Collision resolution strategies: chaining vs. open addressing\n'
        '• Trade-offs: Hash tables trade space for time efficiency'
    )
    
    # Core Techniques
    doc.add_heading('Core Techniques', level=1)
    doc.add_paragraph(
        '• Two-pointer technique (opposite direction, same direction, fast-slow)\n'
        '• Sliding window approach\n'
        '• Prefix sums for range queries\n'
        '• Hash map for frequency counting and lookups\n'
        '• In-place operations to minimize space complexity'
    )
    
    # Essential Problems
    doc.add_heading('Essential Problems', level=1)
    doc.add_paragraph(
        '1. Two Sum: Understand complement mapping and why hash table beats brute force\n'
        '2. Valid Anagram: Character frequency comparison\n'
        '3. Contains Duplicate: Hash set for duplicate detection\n'
        '4. Product of Array Except Self: Left/right prefix products\n'
        '5. Maximum Subarray (Kadane\'s Algorithm): Dynamic programming insight'
    )
    
    # Detailed Problem Explanations
    doc.add_heading('Detailed Problem Explanations', level=1)
    
    doc.add_heading('Two Sum', level=2)
    doc.add_paragraph(
        'Problem: Given an array of integers nums and an integer target, return indices '
        'of the two numbers such that they add up to target.\n\n'
        
        'Approach: Use a hash map to store previously seen values and their indices. '
        'For each element, check if its complement (target - current_value) exists in '
        'the hash map. If so, return the stored index and current index.\n\n'
        
        'Time Complexity: O(n)\n'
        'Space Complexity: O(n)'
    )
    
    doc.add_heading('Contains Duplicate', level=2)
    doc.add_paragraph(
        'Problem: Given an integer array nums, return true if any value appears at '
        'least twice in the array, and return false if every element is distinct.\n\n'
        
        'Approach: Use a hash set to track seen elements. Iterate through the array, '
        'and if an element is already in the set, return true. If the loop completes '
        'without finding duplicates, return false.\n\n'
        
        'Time Complexity: O(n)\n'
        'Space Complexity: O(n)'
    )
    
    doc.add_heading('Product of Array Except Self', level=2)
    doc.add_paragraph(
        'Problem: Given an integer array nums, return an array answer such that '
        'answer[i] is equal to the product of all the elements of nums except nums[i].\n\n'
        
        'Approach: Calculate left and right prefix products. For each position i, '
        'the result is the product of all elements to the left and all elements to '
        'the right. This can be done in two passes without extra space for the result.\n\n'
        
        'Time Complexity: O(n)\n'
        'Space Complexity: O(1) excluding output array'
    )
    
    # Key Insights
    doc.add_heading('Key Insights', level=1)
    doc.add_paragraph(
        '• Hash tables trade space for time efficiency\n'
        '• Array indices represent positions; values represent data\n'
        '• Many problems become tractable with auxiliary space\n'
        '• Two-pointer techniques are powerful for sorted arrays\n'
        '• Prefix computations often simplify range-based queries'
    )
    
    # Practice Tips
    doc.add_heading('Practice Tips', level=1)
    doc.add_paragraph(
        '• Master the two-pointer technique for sorted arrays\n'
        '• Understand when to use hash sets vs hash maps\n'
        '• Practice converting brute force O(n²) solutions to O(n) using hashing\n'
        '• Recognize problems where prefix sums are applicable\n'
        '• Always consider space-time tradeoffs in your solutions'
    )
    
    # Save the document
    filename = 'Day1_Arrays_Hashing.docx'
    filepath = os.path.join(os.getcwd(), filename)
    doc.save(filepath)
    print(f"Document saved as {filepath}")
    
    return filepath

if __name__ == "__main__":
    create_day1_document()